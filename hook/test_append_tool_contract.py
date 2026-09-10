"""追加工具排版契约 hook —— 抓「XML 属性写了、渲染器不认」的排版丢失。

对应踩过的坑：
  1. 首行缩进只写绝对值 w:firstLine，WPS 等中文渲染器优先认 firstLineChars，
     只写绝对值时用户端「没有首行缩进」；
  2. w:shd 用 pPr.append() 手插，排到了 w:spacing 之后 —— CT_PPr 子元素有固定次序
     （shd < spacing < ind），乱序的 pPr 会被严格的渲染器整段丢弃（缩进和底纹一起丢）。
"""
from __future__ import annotations

import importlib.util
import pathlib

import pytest

pytest.importorskip("docx", reason="追加工具排版契约需要 python-docx")

from docx import Document
from docx.oxml.ns import qn

REPO = pathlib.Path(__file__).resolve().parent.parent
TOOL = REPO / "阅读理解题目追加工具" / "阅读理解题目追加工具.py"

SPEC = {
    "type": "阅读单选",
    "passage": ["First paragraph of the passage.", "Second paragraph 你好。"],
    "questions": [{"stem": "1．What is it?", "options": ["A．x", "B．y", "C．z", "D．w"]}],
    "answers": "1．A",
    "summary": "导语文本",
    "details": ["1．详解文本"],
}

# CT_PPr 子元素次序（本工具会用到的子集，按 schema 顺序）
PPR_ORDER = ["pStyle", "shd", "spacing", "ind", "jc", "rPr", "sectPr"]


def _append_to_tmp(tmp_path):
    spec_mod = importlib.util.spec_from_file_location("append_tool", TOOL)
    mod = importlib.util.module_from_spec(spec_mod)
    spec_mod.loader.exec_module(mod)

    out = tmp_path / "out.docx"
    doc = Document()
    doc.add_paragraph("base")
    doc.save(str(out))
    mod.append_from_spec(str(out), SPEC)
    return Document(str(out))


def _ppr_seq(p):
    pPr = p._element.pPr
    return [ch.tag.split("}")[1] for ch in (pPr if pPr is not None else [])]


def test_first_line_indent_uses_char_based_twofold(tmp_path):
    """正文首行缩进必须双写 firstLineChars + firstLine（渲染器各取所需）。"""
    doc = _append_to_tmp(tmp_path)
    for p in doc.paragraphs:
        if p.text.strip().startswith("First paragraph"):
            ind = p._element.pPr.find(qn("w:ind"))
            assert ind is not None, "正文段缺 w:ind"
            assert ind.get(qn("w:firstLineChars")) == "200", "缺字符制缩进（WPS 端会没有首行缩进）"
            assert ind.get(qn("w:firstLine")) is not None, "缺绝对值缩进回退"
            break
    else:
        pytest.fail("没找到追加的正文段")


def test_shd_and_ind_are_in_schema_order(tmp_path):
    """w:shd 必须排在 w:spacing 之前、w:ind 紧跟 spacing —— 乱序会被整段丢弃。"""
    doc = _append_to_tmp(tmp_path)
    shaded = 0
    for p in doc.paragraphs:
        seq = _ppr_seq(p)
        for tag in seq:
            assert tag in PPR_ORDER, f"pPr 出现序列表之外的子元素 {tag}，请核对 CT_PPr 次序"
        idx = [PPR_ORDER.index(t) for t in seq]
        assert idx == sorted(idx), f"pPr 子元素乱序：{seq}"
        if seq and "shd" in seq:
            shaded += 1
    assert shaded == 3, f"答案块应为 答案/导语/详解 共 3 段底纹，实际 {shaded} 段"
